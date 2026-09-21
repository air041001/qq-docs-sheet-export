#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把导出的 CSV 排版成好看的 xlsx（表头样式 / 斑马纹 / 边框 / 列宽 / 冻结 / 筛选）。

用法:
    python format_excel.py <input.csv> [output.xlsx]

也可被 qq_sheet_export.py --xlsx 调用：那里用的是 csv_to_xlsx()。
依赖: openpyxl
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', start_color='4472C4', end_color='4472C4')
HEADER_FONT = Font(name='微软雅黑', bold=True, size=10, color='FFFFFF')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
CELL_FONT = Font(name='微软雅黑', size=9)
CELL_ALIGN = Alignment(vertical='center', wrap_text=True)
ZEBRA_FILL = PatternFill('solid', start_color='F2F7FB', end_color='F2F7FB')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAX_WIDTH = 45
MIN_WIDTH = 6


def clean_cell(value):
    """去掉单元格内多余的空行/首尾空白，保留真实换行。"""
    if value is None:
        return ''
    text = str(value).replace('\r\n', '\n').replace('\r', '\n')
    lines = [line.strip() for line in text.split('\n')]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return '\n'.join(lines)


def calc_width(value):
    """按内容估算列宽（中文按 2 字宽）"""
    if not value:
        return MIN_WIDTH
    lines = str(value).split('\n')
    widest = max((len(line) for line in lines), default=0)
    cjk = sum(1 for ch in str(value) if '\u2e80' <= ch <= '\u9fff' or '\uff00' <= ch <= '\uffef')
    ascii_count = max(widest - cjk, 0)
    return min(max(cjk * 2.1 + ascii_count * 1.1 + 2, MIN_WIDTH), MAX_WIDTH)


def read_csv(path):
    for encoding in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            with open(path, newline='', encoding=encoding) as fh:
                return list(csv.reader(fh))
        except UnicodeDecodeError:
            continue
    raise RuntimeError('无法识别 %s 的编码（试过 utf-8-sig / utf-8 / gbk）' % path)


def csv_to_xlsx(csv_path, xlsx_path=None, sheet_title=None):
    rows = read_csv(csv_path)
    if not rows:
        raise RuntimeError('%s 是空文件' % csv_path)
    xlsx_path = xlsx_path or (os.path.splitext(csv_path)[0] + '.xlsx')

    width = max(len(r) for r in rows)
    header = [clean_cell(h).replace('\n', ' ') for h in rows[0]]
    header += [''] * (width - len(header))
    body = []
    for row in rows[1:]:
        cells = [clean_cell(c) for c in row] + [''] * (width - len(row))
        if any(cells):
            body.append(cells)

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or 'Sheet1')[:31]

    widths = [calc_width(h) for h in header]
    for col, text in enumerate(header, 1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    for r, cells in enumerate(body, 2):
        zebra = (r % 2 == 0)
        for col, text in enumerate(cells, 1):
            cell = ws.cell(row=r, column=col, value=text)
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGN
            cell.border = BORDER
            if zebra:
                cell.fill = ZEBRA_FILL
            widths[col - 1] = max(widths[col - 1], calc_width(text))

    for col, width_value in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width_value + 1
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    return xlsx_path, len(body), width


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        print(__doc__)
        return 2
    src = argv[0]
    dst = argv[1] if len(argv) > 1 else None
    path, nrows, ncols = csv_to_xlsx(src, dst)
    print('%s ← %d 行 × %d 列' % (path, nrows, ncols))
    return 0


if __name__ == '__main__':
    sys.exit(main())
